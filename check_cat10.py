"""
check_cat10.py
"""
import openpyxl

wb = openpyxl.load_workbook('OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx')
master_ws = wb['Master - All Live Verified']

cat_col = None
for c in range(1, master_ws.max_column + 1):
    if master_ws.cell(1, c).value == 'Primary Target Category':
        cat_col = c
        break

print(f"Primary Target Category is at Column {cat_col}")

categories = set()
count = 0
for r in range(2, master_ws.max_row + 1):
    val = str(master_ws.cell(r, cat_col).value or '').strip()
    categories.add(val)
    if 'multiphysics' in val.lower() or 'thermal' in val.lower() or '10.' in val:
        count += 1

print("Matching categories found:")
for cat in sorted(categories):
    if 'multiphysics' in cat.lower() or 'thermal' in cat.lower() or '10.' in cat:
        print(f"  Found: '{cat}'")

print(f"Total matching professors in Master: {count}")
