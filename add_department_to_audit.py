"""
add_department_to_audit.py
==========================
Adds 'Academic Department' to:
  - faculty_status_audit.csv
  - faculty_status_audit.xlsx

Extracted directly by S.N. from 'Master - All Live Verified'.
Maintains:
  - Strict ascending order 1 to 1269
  - Full-row color coding (1=Green, 2=Yellow, 3=Red)
  - Direct clickable hyperlinks on Profile URL
"""

import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_MASTER = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_CFD_AERO_COMPMATH.xlsx")
CSV_PATH = Path(r"d:\Others\findprofs\faculty_status_audit.csv")
XLSX_PATH = Path(r"d:\Others\findprofs\faculty_status_audit.xlsx")

def main():
    print(f"Loading Master workbook to get Academic Departments: {EXCEL_MASTER}")
    wb_master = openpyxl.load_workbook(EXCEL_MASTER, read_only=True)
    ws_master = wb_master["Master - All Live Verified"]

    # Map S.N. -> Academic Department (Col 1: S.N., Col 7: Academic Department)
    dept_map = {}
    for row in ws_master.iter_rows(min_row=2, values_only=True):
        sn = row[0]
        dept = row[6]
        if sn is not None:
            dept_map[str(sn)] = str(dept or "").strip()

    print(f"Mapped {len(dept_map)} faculty departments from Master.")

    # Read current audit CSV
    print(f"Reading current audit CSV: {CSV_PATH}")
    with open(CSV_PATH, "r", encoding="utf-8-sig", errors="ignore") as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = list(reader)

    # Current header:
    # ['S.N.', 'Name', 'University', 'OpenAlex ID', 'Detected Title', 'Web Status Flag', 'Last Active Year', 'Papers (2022-2026)', 'Recommended Action', 'Activeness Rating', 'Profile URL']
    # Insert 'Academic Department' right after 'University' (col index 3)
    new_header = header[:3] + ["Academic Department"] + header[3:]
    print("New Header:", new_header)

    updated_rows = []
    for r in rows:
        sn = str(r[0]).strip()
        dept = dept_map.get(sn, "Unknown Department")

        # Clean URL if it has =HYPERLINK formula
        url_val = r[-1]
        raw_url = url_val
        if 'HYPERLINK("' in url_val:
            parts = url_val.split('"')
            if len(parts) >= 2:
                raw_url = parts[1]

        # Construct new row with Academic Department inserted at pos 3
        # r[:3] -> S.N., Name, University
        # r[3:-1] -> OpenAlex ID, Detected Title, Web Status Flag, Last Active Year, Papers, Recommended Action, Rating
        new_row = r[:3] + [dept] + r[3:-1] + [raw_url]
        updated_rows.append(new_row)

    # Sort strictly by S.N. ascending
    updated_rows.sort(key=lambda r: int(r[0]) if str(r[0]).strip().isdigit() else 999999)

    # 1. Write updated CSV
    with open(CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(new_header)
        for r in updated_rows:
            url = str(r[-1]).strip()
            url_formula = f'=HYPERLINK("{url}","{url}")' if url.startswith("http") else url
            writer.writerow(r[:-1] + [url_formula])

    print(f"[DONE] Saved updated CSV: {CSV_PATH}")

    # 2. Write updated full-row colored Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faculty Status Audit"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    link_font = Font(name="Calibri", size=10, color="0000FF", underline="single")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    fill_rating_1 = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Green
    fill_rating_2 = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Yellow
    fill_rating_3 = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Red

    # Write headers
    for c_idx, h in enumerate(new_header, start=1):
        cell = ws.cell(1, c_idx, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Write data
    for r_idx, r in enumerate(updated_rows, start=2):
        rating = int(r[10]) # Activeness Rating is now at index 10 (11th column)
        if rating == 1:
            row_fill = fill_rating_1
        elif rating == 2:
            row_fill = fill_rating_2
        else:
            row_fill = fill_rating_3

        for c_idx, val in enumerate(r, start=1):
            cell = ws.cell(r_idx, c_idx)
            cell.font = regular_font
            cell.fill = row_fill
            cell.border = thin_border

            # S.N.
            if c_idx == 1:
                cell.value = int(val) if str(val).isdigit() else val
                cell.alignment = center_align
            # Papers count
            elif c_idx == 9:
                cell.value = int(val) if str(val).isdigit() else val
                cell.alignment = center_align
            # Activeness Rating
            elif c_idx == 11:
                cell.value = int(val)
                cell.font = bold_font
                cell.alignment = center_align
            # Profile URL
            elif c_idx == 12:
                url = str(val).strip()
                cell.value = url
                if url.startswith("http"):
                    cell.hyperlink = url
                    cell.font = link_font
                cell.alignment = left_align
            else:
                cell.value = val
                cell.alignment = center_align if c_idx in [5, 7, 8] else left_align

    # Freeze panes & AutoFilter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(new_header))}{len(updated_rows)+1}"

    # Auto-adjust column widths
    widths = {
        "A": 8,   # S.N.
        "B": 28,  # Name
        "C": 34,  # University
        "D": 38,  # Academic Department
        "E": 16,  # OpenAlex ID
        "F": 24,  # Detected Title
        "G": 22,  # Web Status Flag
        "H": 16,  # Last Active Year
        "I": 18,  # Papers (2022-2026)
        "J": 35,  # Recommended Action
        "K": 18,  # Activeness Rating
        "L": 55,  # Profile URL
    }
    for col_let, w in widths.items():
        ws.column_dimensions[col_let].width = w

    wb.save(XLSX_PATH)
    print(f"[DONE] Saved updated full-row colored Excel audit file: {XLSX_PATH}")

if __name__ == "__main__":
    main()
