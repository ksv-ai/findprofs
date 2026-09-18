import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import os

excel_path = r"d:\Others\findprofs\universities_wise\arizona_state_university\asu_aerospace_mechanical_faculty.xlsx"
csv_all = r"d:\Others\findprofs\universities_wise\arizona_state_university\asu_aerospace_mechanical_faculty.csv"
csv_matched = r"d:\Others\findprofs\universities_wise\arizona_state_university\asu_aerospace_mechanical_faculty_field_matched.csv"

df_all = pd.read_csv(csv_all)
df_matched = pd.read_csv(csv_matched)

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

sheets_data = [
    ("Field Matched Faculty (A-Z)", df_matched),
    ("All Active Faculty (A-Z)", df_all)
]

for sheet_title, df in sheets_data:
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    headers = list(df.columns)
    ws.append(headers)

    # Style header row
    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    prof_col_idx = headers.index("Profile URL") if "Profile URL" in headers else None
    scholar_col_idx = headers.index("Google Scholar URL") if "Google Scholar URL" in headers else None
    email_col_idx = headers.index("Email") if "Email" in headers else None

    for r_idx, row in df.iterrows():
        row_values = list(row)
        ws_row = r_idx + 2
        
        # Replace URLs with HYPERLINK formulas so Excel recognizes and displays them as direct clickable links everywhere
        if prof_col_idx is not None:
            url = str(row_values[prof_col_idx])
            if url.startswith("http"):
                row_values[prof_col_idx] = f'=HYPERLINK("{url}", "{url}")'
                
        if scholar_col_idx is not None:
            url = str(row_values[scholar_col_idx])
            if url.startswith("http"):
                name = str(row["Name"])
                row_values[scholar_col_idx] = f'=HYPERLINK("{url}", "Google Scholar Profile")'
                
        if email_col_idx is not None:
            email = str(row_values[email_col_idx])
            if "@" in email:
                row_values[email_col_idx] = f'=HYPERLINK("mailto:{email}", "{email}")'
                
        ws.append(row_values)

        # Style data cells and explicit hyperlinks
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=ws_row, column=c_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Check if this cell is a link
            if c_idx - 1 in [prof_col_idx, scholar_col_idx, email_col_idx]:
                raw_v = str(df.iloc[r_idx, c_idx - 1])
                if raw_v.startswith("http"):
                    cell.hyperlink = raw_v
                    cell.font = link_font
                elif "@" in raw_v:
                    cell.hyperlink = f"mailto:{raw_v}"
                    cell.font = link_font

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        col_name = str(col[0].value)
        for cell in col:
            # handle formula length estimation
            val = str(cell.value or '')
            if val.startswith('='):
                # approximate display length
                if 'Google Scholar Profile' in val:
                    l = 22
                elif 'mailto:' in val:
                    l = len(val.split('"')[3]) if len(val.split('"')) > 3 else 20
                else:
                    l = 40
            else:
                l = len(val)
            if l > max_len:
                max_len = l
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

wb.save(excel_path)
print(f"Successfully saved enhanced clickable Excel file: {excel_path}")
