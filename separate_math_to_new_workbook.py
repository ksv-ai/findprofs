"""
separate_math_to_new_workbook.py
================================
1. Creates a brand-new standalone Excel workbook dedicated exclusively to Mathematics:
   'OpenalexID_R1_FACULTY_MATHEMATICS_COMPMATH.xlsx'
   - Contains all faculty from 'Mathematics & Comp Math'
   - S.N. re-numbered 1 to N
   - Top row frozen (A2), AutoFilter enabled, proper column widths, full styling and colors preserved.

2. From 'OpenalexID_R1_FACULTY_RESEARCH_CFD_AERO_COMPMATH.xlsx':
   - Removes the 'Mathematics & Comp Math' worksheet tab entirely.
   - Removes all purely Math faculty from 'Master - All Live Verified'
     (faculty cross-appointed in Mech or Aero remain in Master).
   - Re-numbers S.N. sequentially (1 to N) in 'Master - All Live Verified'.
   - Refreshes AutoFilter and formatting.
"""

from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

ORIGINAL_EXCEL = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_CFD_AERO_COMPMATH.xlsx")
MATH_EXCEL = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_MATHEMATICS_COMPMATH.xlsx")

def copy_cell(src, dst):
    dst.value = src.value
    if src.font: dst.font = copy(src.font)
    if src.fill: dst.fill = copy(src.fill)
    if src.border: dst.border = copy(src.border)
    if src.alignment: dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format

def format_sheet_properly(ws):
    ws.freeze_panes = "A2"
    max_col_letter = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{max_col_letter}{ws.max_row}"

def main():
    print(f"Loading original workbook: {ORIGINAL_EXCEL}")
    wb_orig = openpyxl.load_workbook(ORIGINAL_EXCEL)

    math_tab = wb_orig["Mathematics & Comp Math"]
    master_tab = wb_orig["Master - All Live Verified"]
    mech_tab = wb_orig["Mechanical Engineering"]
    aero_tab = wb_orig["Aerospace Engineering"]

    cols = master_tab.max_column

    # 1. Gather Engineering keys (Mech or Aero): (clean_name, clean_uni)
    engineering_keys = set()
    for sheet in [mech_tab, aero_tab]:
        for r in range(2, sheet.max_row + 1):
            name = str(sheet.cell(r, 2).value or "").strip().lower()
            uni = str(sheet.cell(r, 3).value or "").strip().lower()
            if name and uni:
                engineering_keys.add((name, uni))

    print(f"Total faculty keys in Mechanical & Aerospace tabs: {len(engineering_keys)}")

    # 2. Gather pure Math keys to remove from Master
    math_keys = set()
    for r in range(2, math_tab.max_row + 1):
        name = str(math_tab.cell(r, 2).value or "").strip().lower()
        uni = str(math_tab.cell(r, 3).value or "").strip().lower()
        if name and uni:
            math_keys.add((name, uni))

    pure_math_keys = {k for k in math_keys if k not in engineering_keys}
    print(f"Total faculty in Math tab: {len(math_keys)}")
    print(f"Pure Math faculty to remove from Master (not in Mech or Aero): {len(pure_math_keys)}")

    # 3. Create the standalone Mathematics workbook
    print(f"\nCreating standalone Mathematics workbook: {MATH_EXCEL}")
    wb_math = openpyxl.Workbook()
    ws_math = wb_math.active
    ws_math.title = "Mathematics & Comp Math"

    # Copy header from math tab
    for c in range(1, cols + 1):
        copy_cell(math_tab.cell(1, c), ws_math.cell(1, c))

    # Copy data rows and renumber S.N. sequentially
    for r_idx, src_r in enumerate(range(2, math_tab.max_row + 1), start=2):
        sn = r_idx - 1
        ws_math.cell(r_idx, 1).value = sn
        ws_math.cell(r_idx, 1).alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
        if math_tab.cell(src_r, 1).font: ws_math.cell(r_idx, 1).font = copy(math_tab.cell(src_r, 1).font)

        for c in range(2, cols + 1):
            copy_cell(math_tab.cell(src_r, c), ws_math.cell(r_idx, c))

    # Set column widths & formatting
    format_sheet_properly(ws_math)
    for col in math_tab.columns:
        col_letter = get_column_letter(col[0].column)
        ws_math.column_dimensions[col_letter].width = math_tab.column_dimensions[col_letter].width or 20

    wb_math.save(MATH_EXCEL)
    print(f"[DONE] Saved standalone Math workbook: {MATH_EXCEL} ({ws_math.max_row - 1} faculty)")

    # 4. Remove purely Math faculty from Master
    print(f"\nPurging pure Math faculty from Master tab...")
    rows_to_delete = []
    for r in range(2, master_tab.max_row + 1):
        name = str(master_tab.cell(r, 2).value or "").strip().lower()
        uni = str(master_tab.cell(r, 3).value or "").strip().lower()
        if (name, uni) in pure_math_keys:
            rows_to_delete.append(r)

    print(f"Rows to delete from Master: {len(rows_to_delete)}")
    for r in reversed(rows_to_delete):
        master_tab.delete_rows(r, 1)

    # Renumber S.N. sequentially in Master (1 to N)
    for idx, r in enumerate(range(2, master_tab.max_row + 1), start=1):
        master_tab.cell(r, 1).value = idx

    format_sheet_properly(master_tab)
    print(f"Master remaining faculty: {master_tab.max_row - 1}")

    # 5. Remove 'Mathematics & Comp Math' worksheet tab from original workbook
    if "Mathematics & Comp Math" in wb_orig.sheetnames:
        del wb_orig["Mathematics & Comp Math"]
        print("Removed 'Mathematics & Comp Math' tab from main workbook.")

    wb_orig.save(ORIGINAL_EXCEL)
    print(f"[DONE] Saved updated main workbook: {ORIGINAL_EXCEL}")

if __name__ == "__main__":
    main()
