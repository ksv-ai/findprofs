import csv
import io
import json
import os
import re
import sys
import time
import requests
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

API_KEY = "NzjuLll4FIEV5HFS2mCK4g"
HEADERS = {
    "Authorization": f"Bearer {API_KEY}",
    "User-Agent": "FindProfs/1.0 (mailto:ksv.ai.research@gmail.com)",
}

EXCEL_PATH = r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx"

# Core target topics in OpenAlex for Aero / Fluids / CFD / Propulsion / Turbulence
TARGET_TOPIC_NAMES = [
    "fluid dynamics and turbulent flows",
    "computational fluid dynamics and aerodynamics",
    "combustion and flame dynamics",
    "rocket and propulsion systems research",
    "combustion and detonation processes",
    "wind and air flow studies",
    "aerodynamics and acoustics in jet flows",
    "gas dynamics and kinetic theory",
    "fluid dynamics simulations and interactions",
    "turbomachinery aerodynamics and aeromechanics",
    "particle dynamics in fluid flows",
    "spacecraft and cryogenic technologies"
]

def load_existing_faculty(uni_substring: str):
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Master - All Live Verified"]
    existing = set()
    existing_list = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, 2).value or "").strip()
        uni = str(ws.cell(r, 3).value or "").strip()
        oa_id = str(ws.cell(r, 6).value or "").strip()
        if uni_substring.lower() in uni.lower():
            clean_name = re.sub(r"[^a-zA-Z]", "", name).lower()
            existing.add(clean_name)
            existing_list.append((name, oa_id))
    return existing, existing_list

def discover_university_aero_faculty(inst_id: str, uni_name: str, min_works=15, min_cites=100):
    print("=" * 80)
    print(f"RUNNING OPENALEX FACULTY DISCOVERY FOR: {uni_name} ({inst_id})")
    print("=" * 80)

    existing_set, existing_list = load_existing_faculty(uni_name)
    print(f"Current faculty recorded in your workbook: {len(existing_list)}")

    candidates = {}
    page = 1
    per_page = 50

    # Query authors currently affiliated with this institution
    # Sorting by works_count desc to capture established faculty & PIs
    while page <= 4:
        url = f"https://api.openalex.org/authors?filter=last_known_institutions.id:{inst_id},works_count:>{min_works}&sort=works_count:desc&per_page={per_page}&page={page}"
        try:
            r = requests.get(url, headers=HEADERS, timeout=15)
            if r.status_code != 200:
                print(f"API returned status {r.status_code}")
                break
            results = r.json().get("results", [])
            if not results:
                break
        except Exception as e:
            print(f"Request error: {e}")
            break

        for author in results:
            author_id = author.get("id", "").split("/")[-1]
            display_name = author.get("display_name", "")
            works_count = author.get("works_count", 0)
            cited_by_count = author.get("cited_by_count", 0)

            # Check topics
            topics = [t.get("display_name", "") for t in (author.get("topics") or [])[:5] if t.get("display_name")]
            matched_topics = [t for t in topics if any(k in t.lower() for k in [
                "fluid", "aerodynamic", "cfd", "turbulen", "combustion", "propulsion", "flame", "aeroacoustic", "hypersonic", "gas dynamics"
            ])]

            if matched_topics and cited_by_count >= min_cites:
                candidates[author_id] = {
                    "author_id": author_id,
                    "name": display_name,
                    "works_count": works_count,
                    "cited_by_count": cited_by_count,
                    "h_index": author.get("summary_stats", {}).get("h_index", 0),
                    "topics": topics[:3],
                    "matched_topic": matched_topics[0]
                }

        page += 1
        time.sleep(0.1)

    print(f"Discovered {len(candidates)} high-impact Aero/Fluids/Propulsion faculty active at {uni_name}!")

    # Diff against current Excel list
    already_have = []
    missing_new = []

    for oa_id, cand in candidates.items():
        clean_cand = re.sub(r"[^a-zA-Z]", "", cand["name"]).lower()
        # Fuzzy / substring check
        is_known = any(clean_cand in ex or ex in clean_cand for ex in existing_set)
        if is_known:
            already_have.append(cand)
        else:
            missing_new.append(cand)

    print(f"\nAlready in your Excel: {len(already_have)}")
    print(f"MISSING FROM YOUR EXCEL (NEW DISCOVERIES): {len(missing_new)}")

    print("\n--- SAMPLE MISSING ACTIVE AERO / CFD / PROPULSION PROFESSORS ---")
    for i, m in enumerate(missing_new[:12], 1):
        print(f"{i:2d}. {m['name']:<28} | ID: {m['author_id']} | Works: {m['works_count']:3d} | Cites: {m['cited_by_count']:5d} | h: {m['h_index']:2d}")
        print(f"    Primary Topic: {m['matched_topic']}")
        print(f"    All Topics:    {m['topics']}")

    # Export report
    out_csv = rf"d:\Others\findprofs\discovered_missing_{re.sub(r'[^a-zA-Z0-9]+', '_', uni_name.lower())}.csv"
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Name", "OpenAlex ID", "Works Count", "Citations", "h-index", "Matched Domain Topic", "Top Topics"])
        for m in missing_new:
            writer.writerow([m["name"], m["author_id"], m["works_count"], m["cited_by_count"], m["h_index"], m["matched_topic"], " | ".join(m["topics"])])
    print(f"\nFull list of missing faculty exported to: {out_csv}")

if __name__ == "__main__":
    # Test on Purdue University
    discover_university_aero_faculty("I219193219", "Purdue", min_works=15, min_cites=150)
