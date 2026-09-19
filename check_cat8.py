"""
check_cat8.py
"""
import openpyxl

wb = openpyxl.load_workbook('OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx')

cat_col = 4 # Primary Target Category is column 4

mech_ws = wb['Mechanical Engineering']
aero_ws = wb['Aerospace Engineering']
math_ws = wb['Mathematics & Comp Math']

def get_cat8_count(ws):
    count = 0
    profs = []
    for r in range(2, ws.max_row + 1):
        cat = str(ws.cell(r, cat_col).value or '').strip()
        if '8.' in cat or 'numerical methods' in cat.lower():
            count += 1
            name = str(ws.cell(r, 1).value or '').strip()
            uni = str(ws.cell(r, 2).value or '').strip()
            profs.append((name, uni))
    return count, profs

mech_count, mech_profs = get_cat8_count(mech_ws)
aero_count, aero_profs = get_cat8_count(aero_ws)
math_count, math_profs = get_cat8_count(math_ws)

print(f"Mechanical Engineering Cat 8 faculty: {mech_count}")
print(f"Aerospace Engineering Cat 8 faculty: {aero_count}")
print(f"Current Mathematics & Comp Math faculty: {math_ws.max_row - 1} (of which Cat 8: {math_count})")
print(f"Total faculty to transfer to Math: {mech_count + aero_count}")

# Check overlap between Mech/Aero Cat 8 and Math existing
math_existing_keys = {
    (str(math_ws.cell(r, 1).value or '').strip().lower(), str(math_ws.cell(r, 2).value or '').strip().lower())
    for r in range(2, math_ws.max_row + 1)
}

overlap_mech = sum(1 for p in mech_profs if (p[0].lower(), p[1].lower()) in math_existing_keys)
overlap_aero = sum(1 for p in aero_profs if (p[0].lower(), p[1].lower()) in math_existing_keys)
print(f"Overlap with existing Math tab: Mech={overlap_mech}, Aero={overlap_aero}")
