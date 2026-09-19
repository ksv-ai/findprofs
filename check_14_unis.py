"""
check_14_unis.py
"""
import openpyxl

NEW_TARGET_UNIS = [
    "Baylor University",
    "Virginia Commonwealth University",
    "Case Western Reserve University",
    "Columbia University",
    "Georgia Institute of Technology",
    "Louisiana State University",
    "Northeastern University",
    "Northwestern University",
    "Old Dominion University",
    "Southern Illinois University Carbondale",
    "Southern Methodist University",
    "Stevens Institute of Technology",
    "University of California, Merced",
    "University of Denver",
]

wb = openpyxl.load_workbook('OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx')
master_ws = wb['Master - All Live Verified']

all_master_unis = {str(master_ws.cell(r, 2).value or '').strip() for r in range(2, master_ws.max_row + 1)}

print("Matching check in Master:")
total_found = 0
for target in NEW_TARGET_UNIS:
    target_clean = target.strip().lower()
    matches = [u for u in all_master_unis if u.lower() == target_clean]
    count = sum(1 for r in range(2, master_ws.max_row + 1) if str(master_ws.cell(r, 2).value or '').strip().lower() == target_clean)
    total_found += count
    if matches:
        print(f"  [EXACT MATCH] {target}: {count} faculty")
    else:
        # Check fuzzy
        near = [u for u in all_master_unis if target_clean in u.lower() or u.lower() in target_clean]
        print(f"  [CHECK FUZZY] {target} -> found in sheet: {near}")

print(f"\nTotal faculty to move from Master: {total_found}")
