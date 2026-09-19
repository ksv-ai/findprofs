"""
harvest_openalex_mechaero_json.py
==================================
Harvests high-fidelity faculty research JSON files directly from OpenAlex API
for all faculty in 'Master - All Live Verified' of OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx.

Output structure:
  d:\\Others\\findprofs\\openalex_mechaero_faculty_json\\universities_wise\\<university_slug>\\<prof_slug>.json
"""
import os
import sys
import io
import re
import json
import time
import logging
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading
import requests
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

EXCEL_PATH = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx")
SHEET_NAME = "Master - All Live Verified"
BASE_OUTPUT_DIR = Path(r"d:\Others\findprofs\openalex_mechaero_faculty_json\universities_wise")
LOG_DIR = Path(r"d:\Others\findprofs\logs")
LOG_DIR.mkdir(parents=True, exist_ok=True)
LOG_FILE = LOG_DIR / "harvest_openalex_mechaero_progress.log"

API_KEY = "NzjuLll4FIEV5HFS2mCK4g"
HEADERS = {
    "Authorization": f"Bearer {API_KEY}",
    "User-Agent": "FindProfs/1.0 (mailto:ksv.ai.research@gmail.com)"
}

NUM_WORKERS = 4
DATE_FROM = "2023-01-01"
DATE_TO = "2026-06-30"

logger = logging.getLogger("openalex_harvest")
logger.setLevel(logging.INFO)
formatter = logging.Formatter("[%(asctime)s] [%(levelname)s] %(message)s", datefmt="%Y-%m-%d %H:%M:%S")

fh = logging.FileHandler(LOG_FILE, encoding="utf-8", mode="a")
fh.setFormatter(formatter)
logger.addHandler(fh)

sh = logging.StreamHandler(sys.stdout)
sh.setFormatter(formatter)
logger.addHandler(sh)

stats_lock = threading.Lock()
processed_count = 0
cached_count = 0
success_count = 0
error_count = 0
total_items = 0
start_time = time.time()

thread_local = threading.local()

def get_session():
    if not hasattr(thread_local, "session"):
        s = requests.Session()
        s.headers.update({
            "Authorization": f"Bearer {API_KEY}",
            "User-Agent": "FindProfs/1.0 (mailto:ksv.ai.research@gmail.com)"
        })
        adapter = requests.adapters.HTTPAdapter(pool_connections=10, pool_maxsize=10, max_retries=2)
        s.mount("https://", adapter)
        s.mount("http://", adapter)
        thread_local.session = s
    return thread_local.session

# Aero & Fluids domain filters
AERO_FLUIDS_KEYWORDS = [
    r"\bcfd\b", r"\bcomputational fluid\b", r"\bfluid dynamic\w*", r"\bfluid mechanic\w*",
    r"\bturbulen\w*", r"\baerodynamic\w*", r"\baeroacoustic\w*", r"\baeroelastic\w*",
    r"\bhypersonic\w*", r"\bsupersonic\w*", r"\btransonic\b", r"\bcompressible flow\b",
    r"\bshock wave\w*", r"\bshock[- ]boundary\b", r"\bsbli\b", r"\bpropulsion\b",
    r"\bcombustion\b", r"\bdetonation\b", r"\brde\b", r"\bscramjet\b", r"\bramjet\b",
    r"\brockets?\b", r"\bnozzle\b", r"\bflame\b", r"\bnavier[- ]stokes\b",
    r"\blarge eddy simulation\b", r"\bles\b", r"\bdirect numerical simulation\b", r"\bdns\b",
    r"\brans\b", r"\bboundary layer\b", r"\bvortex\b", r"\bvortices\b", r"\bvorticity\b",
    r"\bwake\b", r"\bshear flow\b", r"\bmultiphase\b", r"\batomization\b", r"\bdroplet\b",
    r"\bcavitation\b", r"\bparticle[- ]laden\b", r"\bairfoil\b", r"\bwing\b", r"\bpiv\b",
    r"\bparticle image velocimetry\b", r"\bschlieren\b", r"\bwind tunnel\b", r"\bturbomachiner\w*",
    r"\bflight\b", r"\bspacecraft\b", r"\buav\b", r"\bheat transfer\b"
]
COMPILED_AERO = [re.compile(p, re.IGNORECASE) for p in AERO_FLUIDS_KEYWORDS]

EXCLUDE_KEYWORDS = [
    r"\bpatient\w*", r"\bclinical\b", r"\bcancer\b", r"\bcellular\b",
    r"\bbiomedical\b", r"\bcardiovascular\b", r"\bblood\b", r"\bhemodynamic\w*",
    r"\bbone\b", r"\btissue\b", r"\bin vivo\b", r"\bin vitro\b", r"\bsurgery\b",
    r"\bdrug delivery\b", r"\bbio[- ]\w+", r"\bnanocomposite\b"
]
COMPILED_EXCLUDE = [re.compile(p, re.IGNORECASE) for p in EXCLUDE_KEYWORDS]

def score_paper_relevance(title: str, abstract: str, concepts: list) -> int:
    text = f"{title} {abstract} {' '.join(concepts)}"
    exclude_hits = sum(1 for p in COMPILED_EXCLUDE if p.search(text))
    if exclude_hits >= 2:
        return -10
    aero_hits = sum(1 for p in COMPILED_AERO if p.search(text))
    return aero_hits

def slugify(text: str) -> str:
    if not text:
        return "unknown"
    s = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower()).strip("_")
    return s or "unknown"

def reconstruct_abstract(inv_idx: dict) -> str:
    if not inv_idx or not isinstance(inv_idx, dict):
        return ""
    word_positions = []
    for word, positions in inv_idx.items():
        for pos in positions:
            word_positions.append((pos, word))
    word_positions.sort(key=lambda x: x[0])
    abstract = " ".join([wp[1] for wp in word_positions]).strip()
    if len(abstract) > 1200:
        abstract = abstract[:1197] + "..."
    return abstract

def clean_work_obj(w: dict) -> dict:
    primary_loc = w.get("primary_location") or {}
    source = primary_loc.get("source") or {}
    venue = source.get("display_name", "") or ""
    doi = w.get("doi") or primary_loc.get("landing_page_url") or ""
    concepts = [c.get("display_name", "") for c in (w.get("concepts") or [])[:8] if isinstance(c, dict) and c.get("display_name")]
    authors = [a.get("author", {}).get("display_name", "") for a in (w.get("authorships") or []) if isinstance(a, dict) and a.get("author", {}).get("display_name")]
    oa_info = w.get("open_access") or {}

    return {
        "title": w.get("title") or "",
        "publication_year": w.get("publication_year"),
        "publication_date": w.get("publication_date") or "",
        "doi": doi,
        "venue": venue,
        "type": w.get("type") or "",
        "cited_by_count": w.get("cited_by_count") or 0,
        "is_oa": oa_info.get("is_oa", False),
        "oa_url": oa_info.get("oa_url") or "",
        "concepts": concepts,
        "abstract": reconstruct_abstract(w.get("abstract_inverted_index")),
        "authors": authors,
    }

def api_get_with_retry(url: str, params: dict = None, retries: int = 3) -> dict:
    session = get_session()
    for attempt in range(retries):
        try:
            r = session.get(url, params=params, timeout=(8, 15))
            if r.status_code == 200:
                return r.json()
            elif r.status_code == 429:
                time.sleep(2.0 * (attempt + 1))
            elif r.status_code == 404:
                return None
            else:
                time.sleep(1.0)
        except Exception:
            time.sleep(1.5 * (attempt + 1))
    return None

def process_faculty(faculty_info: dict):
    global processed_count, cached_count, success_count, error_count

    name = faculty_info["name"]
    uni = faculty_info["uni"]
    department = faculty_info["department"]
    oa_id = faculty_info["oa_id"]
    row_idx = faculty_info["row_idx"]

    uni_slug = slugify(uni)
    prof_slug = slugify(name)
    dest_dir = BASE_OUTPUT_DIR / uni_slug
    dest_file = dest_dir / f"{prof_slug}.json"

    if dest_file.exists() and dest_file.stat().st_size > 100:
        with stats_lock:
            processed_count += 1
            cached_count += 1
            curr_done = processed_count
        logger.info(f"[{curr_done:3d}/{total_items:3d}] [ALREADY CACHED] {name} | {uni} -> {dest_file.name}")
        return

    clean_author_id = oa_id.replace("https://openalex.org/", "").strip()
    author_url = f"https://api.openalex.org/authors/{clean_author_id}"
    author_data = api_get_with_retry(author_url)

    if not author_data:
        with stats_lock:
            processed_count += 1
            error_count += 1
            curr_done = processed_count
        logger.error(f"[{curr_done:3d}/{total_items:3d}] [NOT FOUND] Row {row_idx}: {name} (ID: {clean_author_id})")
        return

    summary_stats = author_data.get("summary_stats") or {}
    h_index = summary_stats.get("h_index", 0)
    works_count = author_data.get("works_count", 0)
    cited_by_count = author_data.get("cited_by_count", 0)
    display_name = author_data.get("display_name", name)

    raw_topics = author_data.get("topics") or []
    top_topics = [
        {"topic": t.get("display_name", ""), "count": t.get("count", 0)}
        for t in raw_topics[:5] if isinstance(t, dict) and t.get("display_name")
    ]

    time.sleep(0.08)
    top_cited_url = f"https://api.openalex.org/works?filter=author.id:{clean_author_id}&sort=cited_by_count:desc&per_page=10"
    top_cited_data = api_get_with_retry(top_cited_url)
    raw_top = [clean_work_obj(w) for w in (top_cited_data.get("results") or [])] if top_cited_data else []
    
    # Sort top cited by Aero/Fluids relevance
    raw_top.sort(key=lambda w: score_paper_relevance(w["title"], w["abstract"], w["concepts"]), reverse=True)
    top_cited_works = raw_top[:3]

    time.sleep(0.08)
    # Fetch candidate window of recent works to apply Aero / Fluids filter
    recent_url = f"https://api.openalex.org/works?filter=author.id:{clean_author_id},from_publication_date:{DATE_FROM},to_publication_date:{DATE_TO}&sort=publication_date:desc&per_page=15"
    recent_data = api_get_with_retry(recent_url)
    raw_recent = [clean_work_obj(w) for w in (recent_data.get("results") or [])] if recent_data else []

    # Score and filter recent works
    scored_recent = []
    for w in raw_recent:
        score = score_paper_relevance(w["title"], w["abstract"], w["concepts"])
        scored_recent.append((score, w))

    aero_recent = [w for s, w in scored_recent if s > 0]
    neutral_recent = [w for s, w in scored_recent if s == 0]
    filtered_recent = (aero_recent + neutral_recent)[:5]
    if not filtered_recent:
        filtered_recent = raw_recent[:5]

    faculty_record = {
        "faculty_name": name,
        "full_name": display_name,
        "department": department,
        "university": uni,
        "author_id": clean_author_id,
        "works_count": works_count,
        "cited_by_count": cited_by_count,
        "h_index": h_index,
        "top_topics": top_topics,
        "top_cited_works": top_cited_works,
        "recent_works": filtered_recent
    }

    dest_dir.mkdir(parents=True, exist_ok=True)
    tmp_file = dest_file.with_suffix(".tmp")
    with open(tmp_file, "w", encoding="utf-8") as f:
        json.dump(faculty_record, f, indent=2, ensure_ascii=False)
    tmp_file.replace(dest_file)

    with stats_lock:
        processed_count += 1
        success_count += 1
        curr_done = processed_count
        elapsed = time.time() - start_time
        rate = curr_done / elapsed if elapsed > 0 else 0
        rem_sec = (total_items - curr_done) / rate if rate > 0 else 0

    logger.info(
        f"[{curr_done:3d}/{total_items:3d}] ({curr_done/total_items*100:5.1f}%) [SAVED] {name[:25]:<25} | {uni[:25]:<25} | "
        f"Works: {works_count:3d}, Cites: {cited_by_count:5d}, h-index: {h_index:2d} | "
        f"Speed: {rate:.1f} profs/s | ETA: {rem_sec/60:.1f}m"
    )


def main():
    global total_items, start_time
    logger.info("=" * 80)
    logger.info("STARTING OPENALEX FACULTY RESEARCH JSON EXTRACTION PIPELINE")
    logger.info(f"Target Sheet: {SHEET_NAME}")
    logger.info(f"Destination:  {BASE_OUTPUT_DIR}")
    logger.info(f"Date Filter:  {DATE_FROM} to {DATE_TO} (Mid-2026)")
    logger.info("=" * 80)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    faculty_list = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, 2).value or "").strip()
        uni = str(ws.cell(r, 3).value or "").strip()
        department = str(ws.cell(r, 7).value or "").strip()
        oa_id = str(ws.cell(r, 6).value or "").strip()
        if oa_id and oa_id.startswith("A") and len(oa_id) > 3:
            faculty_list.append({"row_idx": r, "name": name, "uni": uni, "department": department, "oa_id": oa_id})

    total_items = len(faculty_list)
    logger.info(f"Total verified faculty members with valid OpenAlex ID: {total_items}")
    logger.info(f"Running multi-threaded harvesting pool with {NUM_WORKERS} workers...")
    start_time = time.time()

    with ThreadPoolExecutor(max_workers=NUM_WORKERS) as executor:
        futures = [executor.submit(process_faculty, f_info) for f_info in faculty_list]
        for _ in as_completed(futures):
            pass

    elapsed_total = time.time() - start_time
    logger.info("=" * 80)
    logger.info("EXTRACTION PIPELINE COMPLETED")
    logger.info(f"Total Processed: {processed_count}")
    logger.info(f"Successfully Harvested & Saved: {success_count}")
    logger.info(f"Already Cached (Skipped):       {cached_count}")
    logger.info(f"Errors / Not Found:             {error_count}")
    logger.info(f"Total Elapsed Time:             {elapsed_total/60:.2f} minutes")
    logger.info(f"Output Directory:               {BASE_OUTPUT_DIR}")
    logger.info("=" * 80)


if __name__ == "__main__":
    main()
