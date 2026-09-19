"""
move_others_universities.py
===========================
Moves all faculty from the specified 18 universities out of:
  - Master - All Live Verified
  - Aerospace Engineering
  - Mathematics & Comp Math
  - Mechanical Engineering

and places them into a dedicated new sheet:
  - Others University Faculties
"""

from copy import copy
from pathlib import Path
import openpyxl

EXCEL_PATH = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx")

TARGET_UNIS = [
    "University of Alaska Fairbanks",
    "Naval Postgraduate School",
    "Portland State University",
    "Clarkson University",
    "Villanova University",
    "Western Michigan University",
    "University of Akron",
    "Cleveland State University",
    "South Dakota School of Mines and Technology",
    "New Mexico Tech",
    "University of New Orleans",
    "Marquette University",
    "Northern Illinois University",
    "Oakland University",
    "Louisiana Tech University",
    "University of South Alabama",
    "University of Texas Rio Grande Valley",
    "Georgia Southern University",
]

TARGET_UNIS_LOWER = {u.strip().lower() for u in TARGET_UNIS}
NEW_TAB_NAME = "Others University Faculties"

def main():
    print(f"Loading workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    master_ws = wb["Master - All Live Verified"]

    # 1. Create or recreate 'Others University Faculties'
    if NEW_TAB_NAME in wb.sheetnames:
        del wb[NEW_TAB_NAME]
    others_ws = wb.create_sheet(title=NEW_TAB_NAME)

    # Copy header row from Master
    for c in range(1, master_ws.max_column + 1):
        src = master_ws.cell(1, c)
        dst = others_ws.cell(1, c)
        dst.value = src.value
        if src.font: dst.font = copy(src.font)
        if src.fill: dst.fill = copy(src.fill)
        if src.border: dst.border = copy(src.border)
        if src.alignment: dst.alignment = copy(src.alignment)

    # 2. Extract rows belonging to the target universities from Master
    others_row = 2
    for r in range(2, master_ws.max_row + 1):
        uni = str(master_ws.cell(r, 2).value or "").strip().lower()
        if uni in TARGET_UNIS_LOWER:
            for c in range(1, master_ws.max_column + 1):
                src = master_ws.cell(r, c)
                dst = others_ws.cell(others_row, c)
                dst.value = src.value
                if src.font: dst.font = copy(src.font)
                if src.fill: dst.fill = copy(src.fill)
                if src.border: dst.border = copy(src.border)
                if src.alignment: dst.alignment = copy(src.alignment)
                dst.number_format = src.number_format
            others_row += 1

    total_moved = others_row - 2
    print(f"Extracted {total_moved} faculty rows into '{NEW_TAB_NAME}'")

    # 3. Purge those universities from Master and Department tabs
    faculty_sheets = [
        "Master - All Live Verified",
        "Aerospace Engineering",
        "Mathematics & Comp Math",
        "Mechanical Engineering",
    ]

    for sheet_name in faculty_sheets:
        ws = wb[sheet_name]
        kept_rows = [1] # Keep header
        for r in range(2, ws.max_row + 1):
            uni = str(ws.cell(r, 2).value or "").strip().lower()
            if uni not in TARGET_UNIS_LOWER:
                kept_rows.append(r)

        removed_count = ws.max_row - len(kept_rows)
        print(f"[{sheet_name}]: Keeping {len(kept_rows)-1} rows, removed {removed_count} rows")

        # Snapshot kept rows
        kept_data = []
        for r in kept_rows:
            row_cells = []
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                row_cells.append({
                    "value": cell.value,
                    "font": copy(cell.font) if cell.font else None,
                    "fill": copy(cell.fill) if cell.fill else None,
                    "border": copy(cell.border) if cell.border else None,
                    "alignment": copy(cell.alignment) if cell.alignment else None,
                    "number_format": cell.number_format,
                })
            kept_data.append(row_cells)

        # Clear and rewrite
        ws.delete_rows(1, ws.max_row + 10)
        for r_idx, row_cells in enumerate(kept_data, start=1):
            for c_idx, data in enumerate(row_cells, start=1):
                cell = ws.cell(r_idx, c_idx)
                cell.value = data["value"]
                if data["font"]: cell.font = data["font"]
                if data["fill"]: cell.fill = data["fill"]
                if data["border"]: cell.border = data["border"]
                if data["alignment"]: cell.alignment = data["alignment"]
                if data["number_format"]: cell.number_format = data["number_format"]

    # 4. Set final sheet ordering
    desired_order = [
        "Master - All Live Verified",
        "Aerospace Engineering",
        "Mathematics & Comp Math",
        "Mechanical Engineering",
        NEW_TAB_NAME,
        "Target Category Breakdown",
    ]
    wb._sheets = [wb[s] for s in desired_order if s in wb.sheetnames]

    # Save
    wb.save(EXCEL_PATH)
    print("\n[DONE] Successfully moved faculty and saved workbook!")

if __name__ == "__main__":
    main()
