"""
reorder_columns_and_tabs.py
===========================
Utility script to easily reorder, rename, or rearrange columns and sheets
in the R1 Faculty Excel workbook.

You can modify:
  1. DESIRED_COLUMN_ORDER: Change the list order below to rearrange columns.
  2. COLUMN_RENAMES: Rename any column title dynamically.
  3. DESIRED_SHEET_ORDER: Change the worksheet tab ordering.

Run with:
    python reorder_columns_and_tabs.py
"""

from pathlib import Path
from copy import copy
import openpyxl

# ── File configuration ────────────────────────────────────────────────────────
EXCEL_PATH = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_CFD_AERO_COMPMATH.xlsx")

# ── 1. Column Renaming Rule ───────────────────────────────────────────────────
# "Old Header Name": "New Header Name"
COLUMN_RENAMES = {
    "Faculty Full Name": "Name",
    # Add any other renames here if needed in the future:
    # "Verified Email Address": "Email",
}

# ── 2. Desired Column Ordering ────────────────────────────────────────────────
# Set the exact column order you want across the faculty tabs.
# Any column not listed here will be placed at the end in its original order.
DESIRED_COLUMN_ORDER = [
    "Name",
    "University",
    "Verified Research Focus & Technical Expertise",
    "Primary Target Category",
    "OpenAlex ID",
    "Academic Department",
    "Verified Email Address",
    "Official Profile URL",
    "Google Scholar URL",
]

# ── 3. Desired Worksheet Tab Ordering ─────────────────────────────────────────
DESIRED_SHEET_ORDER = [
    "Master - All Live Verified",
    "Aerospace Engineering",
    "Mathematics & Comp Math",
    "Mechanical Engineering",
    "Others University Faculties",
    "Target Category Breakdown",
]

# Tabs to apply column reordering to (skip summary tabs like breakdown)
FACULTY_TABS = [
    "Master - All Live Verified",
    "Aerospace Engineering",
    "Mathematics & Comp Math",
    "Mechanical Engineering",
    "Others University Faculties",
]


def reorder_sheet_columns(ws, desired_order: list[str], renames: dict[str, str]):
    """Reorder and rename columns in a worksheet while preserving cell styles."""
    # 1. Apply renames to row 1
    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(1, col).value or "").strip()
        if val in renames:
            ws.cell(1, col).value = renames[val]

    # 2. Get current headers
    current_headers = [str(ws.cell(1, col).value or "").strip() for col in range(1, ws.max_column + 1)]
    
    # 3. Determine new column order mapping
    new_order = []
    for col_name in desired_order:
        if col_name in current_headers:
            new_order.append(col_name)
    
    # Append any remaining columns that weren't explicitly listed
    for col_name in current_headers:
        if col_name not in new_order and col_name:
            new_order.append(col_name)

    # If the order is already identical, nothing to do
    if current_headers == new_order:
        print(f"  [{ws.title}] Columns already in desired order.")
        return

    print(f"  [{ws.title}] Reordering columns to: {new_order[:4]} ...")

    # Map target position to source column index (1-based)
    src_indices = [current_headers.index(col_name) + 1 for col_name in new_order]

    # Read all rows into memory with their cell attributes
    row_data = []
    for r in range(1, ws.max_row + 1):
        row_vals = []
        for src_col in src_indices:
            cell = ws.cell(r, src_col)
            row_vals.append({
                "value": cell.value,
                "font": copy(cell.font) if cell.font else None,
                "fill": copy(cell.fill) if cell.fill else None,
                "border": copy(cell.border) if cell.border else None,
                "alignment": copy(cell.alignment) if cell.alignment else None,
                "number_format": cell.number_format,
            })
        row_data.append(row_vals)

    # Write back the reordered columns
    for r_idx, row_vals in enumerate(row_data, start=1):
        for c_idx, data in enumerate(row_vals, start=1):
            cell = ws.cell(r_idx, c_idx)
            cell.value = data["value"]
            if data["font"]: cell.font = data["font"]
            if data["fill"]: cell.fill = data["fill"]
            if data["border"]: cell.border = data["border"]
            if data["alignment"]: cell.alignment = data["alignment"]
            if data["number_format"]: cell.number_format = data["number_format"]


def main():
    print(f"Loading workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # 1. Reorder sheets
    print("\n--- Reordering Tabs ---")
    current_sheets = wb.sheetnames
    ordered_sheets = [s for s in DESIRED_SHEET_ORDER if s in current_sheets]
    # Add any remaining sheets not mentioned in list
    for s in current_sheets:
        if s not in ordered_sheets:
            ordered_sheets.append(s)
    wb._sheets = [wb[s] for s in ordered_sheets]
    print("New Sheet Order:", wb.sheetnames)

    # 2. Reorder columns across faculty tabs
    print("\n--- Reordering Columns in Tabs ---")
    for tab_name in FACULTY_TABS:
        if tab_name in wb.sheetnames:
            reorder_sheet_columns(wb[tab_name], DESIRED_COLUMN_ORDER, COLUMN_RENAMES)

    # 3. Save workbook
    wb.save(EXCEL_PATH)
    print(f"\n[DONE] Workbook saved: {EXCEL_PATH}")


if __name__ == "__main__":
    main()
