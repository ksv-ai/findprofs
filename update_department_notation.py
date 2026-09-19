"""
update_department_notation.py
=============================
Replaces the full Academic Department name with short notation:
  - 'Mech' (Mechanical Engineering)
  - 'Aero' (Aerospace Engineering)
  - 'Math' (Mathematics & Comp Math)
  (or 'Mech/Aero', 'Aero/Math' if faculty appears in multiple department tabs)

Applies to both:
  - faculty_status_audit.csv
  - faculty_status_audit.xlsx
"""

import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_MASTER = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_CFD_AERO_COMPMATH.xlsx")
CSV_PATH = Path(r"d:\Others\findprofs\faculty_status_audit.csv")
XLSX_PATH = Path(r"d:\Others\findprofs\faculty_status_audit.xlsx")

def main():
    print(f"Loading Master workbook to map department notations: {EXCEL_MASTER}")
    wb_master = openpyxl.load_workbook(EXCEL_MASTER, read_only=True)

    # Build key sets for each department tab: (clean_name, clean_uni)
    def get_keys(sheet_name):
        ws = wb_master[sheet_name]
        keys = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            # Col 1: S.N., Col 2: Name, Col 3: University
            name = str(row[1] or "").strip().lower()
            uni = str(row[2] or "").strip().lower()
            if name and uni:
                keys.add((name, uni))
        return keys

    mech_keys = get_keys("Mechanical Engineering")
    aero_keys = get_keys("Aerospace Engineering")
    math_keys = get_keys("Mathematics & Comp Math")

    print(f"Loaded department keys: Mech={len(mech_keys)}, Aero={len(aero_keys)}, Math={len(math_keys)}")

    # Read current audit CSV
    print(f"Reading: {CSV_PATH}")
    with open(CSV_PATH, "r", encoding="utf-8-sig", errors="ignore") as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = list(reader)

    # Find department column index
    dept_col_idx = header.index("Academic Department")
    header[dept_col_idx] = "Department" # Clean header name

    updated_rows = []
    notation_counts = {}

    for r in rows:
        name = str(r[1]).strip().lower()
        uni = str(r[2]).strip().lower()
        key = (name, uni)

        tags = []
        if key in mech_keys:
            tags.append("Mech")
        if key in aero_keys:
            tags.append("Aero")
        if key in math_keys:
            tags.append("Math")

        # Fallback to textual department keyword if not in the 3 tabs
        if not tags:
            dept_raw = str(r[dept_col_idx]).lower()
            if "aero" in dept_raw:
                tags.append("Aero")
            elif "mech" in dept_raw:
                tags.append("Mech")
            elif "math" in dept_raw or "computational" in dept_raw:
                tags.append("Math")
            else:
                tags.append("Other")

        notation = "/".join(tags)
        notation_counts[notation] = notation_counts.get(notation, 0) + 1

        # Replace department with short notation
        new_row = list(r)
        new_row[dept_col_idx] = notation

        # Clean URL if it has =HYPERLINK formula
        url_val = new_row[-1]
        raw_url = url_val
        if 'HYPERLINK("' in url_val:
            parts = url_val.split('"')
            if len(parts) >= 2:
                raw_url = parts[1]
        new_row[-1] = raw_url

        updated_rows.append(new_row)

    print("Notation breakdown:")
    for k, v in sorted(notation_counts.items()):
        print(f"  {k}: {v}")

    # 1. Save updated CSV
    with open(CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for r in updated_rows:
            url = str(r[-1]).strip()
            url_formula = f'=HYPERLINK("{url}","{url}")' if url.startswith("http") else url
            writer.writerow(r[:-1] + [url_formula])

    print(f"[DONE] Saved CSV: {CSV_PATH}")

    # 2. Save styled XLSX version
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faculty Status Audit"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    link_font = Font(name="Calibri", size=10, color="0000FF", underline="single")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    fill_rating_1 = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_rating_2 = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_rating_3 = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for c_idx, h in enumerate(header, start=1):
        cell = ws.cell(1, c_idx, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    rating_col_idx = header.index("Activeness Rating") + 1

    for r_idx, r in enumerate(updated_rows, start=2):
        rating = int(r[rating_col_idx - 1])
        if rating == 1:
            row_fill = fill_rating_1
        elif rating == 2:
            row_fill = fill_rating_2
        else:
            row_fill = fill_rating_3

        for c_idx, val in enumerate(r, start=1):
            cell = ws.cell(r_idx, c_idx)
            cell.font = regular_font
            cell.fill = row_fill
            cell.border = thin_border

            # S.N.
            if c_idx == 1:
                cell.value = int(val) if str(val).isdigit() else val
                cell.alignment = center_align
            # Department notation
            elif c_idx == (dept_col_idx + 1):
                cell.value = val
                cell.font = bold_font
                cell.alignment = center_align
            # Papers count
            elif c_idx == 9:
                cell.value = int(val) if str(val).isdigit() else val
                cell.alignment = center_align
            # Rating
            elif c_idx == rating_col_idx:
                cell.value = int(val)
                cell.font = bold_font
                cell.alignment = center_align
            # URL
            elif c_idx == len(r):
                url = str(val).strip()
                cell.value = url
                if url.startswith("http"):
                    cell.hyperlink = url
                    cell.font = link_font
                cell.alignment = left_align
            else:
                cell.value = val
                cell.alignment = center_align if c_idx in [5, 7, 8] else left_align

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(updated_rows)+1}"

    widths = {
        "A": 8,   # S.N.
        "B": 28,  # Name
        "C": 34,  # University
        "D": 14,  # Department (Short notation)
        "E": 16,  # OpenAlex ID
        "F": 24,  # Detected Title
        "G": 22,  # Web Status Flag
        "H": 16,  # Last Active Year
        "I": 18,  # Papers (2022-2026)
        "J": 35,  # Recommended Action
        "K": 18,  # Activeness Rating
        "L": 55,  # Profile URL
    }
    for col_let, w in widths.items():
        ws.column_dimensions[col_let].width = w

    wb.save(XLSX_PATH)
    print(f"[DONE] Saved XLSX: {XLSX_PATH}")

if __name__ == "__main__":
    main()
