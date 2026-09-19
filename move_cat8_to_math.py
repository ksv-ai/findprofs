"""
move_cat8_to_math.py
====================
Moves all faculty with category:
  '8. Numerical Methods & Computational Techniques'
out of:
  - Mechanical Engineering
  - Aerospace Engineering
and places them completely into:
  - Mathematics & Comp Math (avoiding duplicate rows if already present)

Maintains table formatting, freeze panes, autofilters, and column widths.
"""

from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

EXCEL_PATH = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx")
TARGET_CAT = "8. Numerical Methods & Computational Techniques"

def format_sheet_properly(ws):
    ws.freeze_panes = "A2"
    max_col_letter = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{max_col_letter}{ws.max_row}"

def is_cat8(val):
    s = str(val or "").strip()
    return "8." in s or "numerical methods" in s.lower()

def main():
    print(f"Loading workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    mech_ws = wb["Mechanical Engineering"]
    aero_ws = wb["Aerospace Engineering"]
    math_ws = wb["Mathematics & Comp Math"]

    cols = math_ws.max_column
    cat_col = 4 # Primary Target Category

    # 1. Existing Math keys to avoid duplicates: (clean_name, clean_uni)
    math_existing_keys = set()
    for r in range(2, math_ws.max_row + 1):
        name = str(math_ws.cell(r, 1).value or "").strip().lower()
        uni = str(math_ws.cell(r, 2).value or "").strip().lower()
        if name and uni:
            math_existing_keys.add((name, uni))

    print(f"Initial Math faculty count: {math_ws.max_row - 1}")

    # 2. Extract Cat 8 rows from Mechanical Engineering and Aerospace Engineering
    faculty_to_add_to_math = []

    for source_name, ws in [("Mechanical Engineering", mech_ws), ("Aerospace Engineering", aero_ws)]:
        initial_count = ws.max_row - 1
        rows_to_delete = []
        
        for r in range(2, ws.max_row + 1):
            cat_val = ws.cell(r, cat_col).value
            if is_cat8(cat_val):
                rows_to_delete.append(r)
                name = str(ws.cell(r, 1).value or "").strip()
                uni = str(ws.cell(r, 2).value or "").strip()
                key = (name.lower(), uni.lower())

                # If not already present in Math, record row data to append
                if key not in math_existing_keys:
                    math_existing_keys.add(key)
                    row_data = []
                    for c in range(1, cols + 1):
                        cell = ws.cell(r, c)
                        row_data.append((
                            cell.value,
                            copy(cell.fill),
                            copy(cell.font),
                            copy(cell.alignment),
                            cell.number_format
                        ))
                    faculty_to_add_to_math.append(row_data)

        # Delete Cat 8 rows from source sheet (bottom-up)
        for r in reversed(rows_to_delete):
            ws.delete_rows(r, 1)

        print(f"[{source_name}] Removed {len(rows_to_delete)} Cat 8 professors. Remaining faculty: {ws.max_row - 1}")

    print(f"Adding {len(faculty_to_add_to_math)} unique Cat 8 professors into 'Mathematics & Comp Math'...")

    # 3. Append to Mathematics & Comp Math
    current_math_row = math_ws.max_row + 1
    for row_vals in faculty_to_add_to_math:
        for c_idx, (val, fill, font, align, num_fmt) in enumerate(row_vals, start=1):
            cell = math_ws.cell(current_math_row, c_idx)
            cell.value = val
            if fill: cell.fill = fill
            if font: cell.font = font
            if align: cell.alignment = align
            if num_fmt: cell.number_format = num_fmt
        current_math_row += 1

    print(f"New 'Mathematics & Comp Math' faculty count: {math_ws.max_row - 1}")

    # 4. Refresh formatting
    format_sheet_properly(mech_ws)
    format_sheet_properly(aero_ws)
    format_sheet_properly(math_ws)

    # 5. Save workbook
    print("Saving workbook...")
    wb.save(EXCEL_PATH)
    print("[DONE] Successfully moved Category 8 professors completely to Mathematics & Comp Math!")

if __name__ == "__main__":
    main()
