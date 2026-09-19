"""
fast_move_others.py
===================
Quickly moves faculty from the specified 18 universities to 'Others University Faculties'
by building new clean sheets, preserving cell attributes cleanly.
"""

from copy import copy
from pathlib import Path
import openpyxl

EXCEL_PATH = Path(r"d:\Others\findprofs\OpenalexID_R1_FACULTY_RESEARCH_AERO_MECH.xlsx")

TARGET_UNIS = [
    "University of Alaska Fairbanks",
    "Naval Postgraduate School",
    "Portland State University",
    "Clarkson University",
    "Villanova University",
    "Western Michigan University",
    "University of Akron",
    "Cleveland State University",
    "South Dakota School of Mines and Technology",
    "New Mexico Tech",
    "University of New Orleans",
    "Marquette University",
    "Northern Illinois University",
    "Oakland University",
    "Louisiana Tech University",
    "University of South Alabama",
    "University of Texas Rio Grande Valley",
    "Georgia Southern University",
]
TARGET_UNIS_LOWER = {u.strip().lower() for u in TARGET_UNIS}
NEW_TAB_NAME = "Others University Faculties"

def copy_cell(src, dst):
    dst.value = src.value
    if src.font: dst.font = copy(src.font)
    if src.fill: dst.fill = copy(src.fill)
    if src.border: dst.border = copy(src.border)
    if src.alignment: dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format

def main():
    print(f"Loading workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # 1. Gather all rows from Master that belong to Others
    master_ws = wb["Master - All Live Verified"]
    max_cols = master_ws.max_column

    others_rows = []
    for r in range(2, master_ws.max_row + 1):
        uni = str(master_ws.cell(r, 2).value or "").strip().lower()
        if uni in TARGET_UNIS_LOWER:
            others_rows.append([master_ws.cell(r, c) for c in range(1, max_cols + 1)])

    print(f"Total faculty to place into '{NEW_TAB_NAME}': {len(others_rows)}")

    # 2. Re-create Others tab
    if NEW_TAB_NAME in wb.sheetnames:
        del wb[NEW_TAB_NAME]
    others_ws = wb.create_sheet(title=NEW_TAB_NAME)

    # Copy header
    for c in range(1, max_cols + 1):
        copy_cell(master_ws.cell(1, c), others_ws.cell(1, c))

    # Populate Others
    for r_idx, row_cells in enumerate(others_rows, start=2):
        for c_idx, src in enumerate(row_cells, start=1):
            copy_cell(src, others_ws.cell(r_idx, c_idx))

    # 3. Filter the 4 faculty tabs in-place by creating clean replacement sheets
    faculty_sheets = [
        "Master - All Live Verified",
        "Aerospace Engineering",
        "Mathematics & Comp Math",
        "Mechanical Engineering",
    ]

    for s_name in faculty_sheets:
        old_ws = wb[s_name]
        cols = old_ws.max_column
        temp_title = f"{s_name}__temp"
        temp_ws = wb.create_sheet(title=temp_title)

        # Copy header
        for c in range(1, cols + 1):
            copy_cell(old_ws.cell(1, c), temp_ws.cell(1, c))

        # Copy only non-target faculty
        new_row = 2
        for r in range(2, old_ws.max_row + 1):
            uni = str(old_ws.cell(r, 2).value or "").strip().lower()
            if uni not in TARGET_UNIS_LOWER:
                for c in range(1, cols + 1):
                    copy_cell(old_ws.cell(r, c), temp_ws.cell(new_row, c))
                new_row += 1

        kept = new_row - 2
        removed = (old_ws.max_row - 1) - kept
        print(f"  [{s_name}] Kept: {kept}, Removed: {removed}")

        # Replace old sheet with new sheet
        del wb[s_name]
        temp_ws.title = s_name

    # 4. Final sheet order
    desired_order = [
        "Master - All Live Verified",
        "Aerospace Engineering",
        "Mathematics & Comp Math",
        "Mechanical Engineering",
        NEW_TAB_NAME,
        "Target Category Breakdown",
    ]
    wb._sheets = [wb[s] for s in desired_order if s in wb.sheetnames]

    # Save
    print(f"Saving workbook...")
    wb.save(EXCEL_PATH)
    print("[DONE] Successfully separated faculties into Others University Faculties and saved!")

if __name__ == "__main__":
    main()
